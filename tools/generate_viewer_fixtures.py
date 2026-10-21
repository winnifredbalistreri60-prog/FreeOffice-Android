#!/usr/bin/env python3
"""Generate synthetic inputs ONLY. Does not open the app or run any tests."""
from pathlib import Path
import sys, hashlib, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.pdfencrypt import StandardEncryption
from PIL import Image, ImageDraw
out=Path(sys.argv[1] if len(sys.argv)>1 else 'synthetic-fixtures');out.mkdir(parents=True,exist_ok=True)
def save(d,n):d.save(out/n)
def fresh():
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);return d
d=fresh()
for i in range(1,4):d.add_heading('Heading level '+str(i),i);d.add_paragraph('Body paragraph below heading.')
save(d,'headings.docx')
d=fresh()
for name,left,right,first in [('left',2,0,0),('right',0,2,0),('first',0,0,1),('hanging',2,0,-1)]:
 p=d.add_paragraph((name+' indentation: visible sample line. ')*12);f=p.paragraph_format;f.left_indent=Cm(left);f.right_indent=Cm(right);f.first_line_indent=Cm(first)
save(d,'indents.docx')
d=fresh()
for a in [WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.RIGHT,WD_ALIGN_PARAGRAPH.JUSTIFY]:p=d.add_paragraph(('Alignment '+str(a)+' sample text. ')*15);p.alignment=a
save(d,'align.docx')
d=fresh()
for line in [1.0,1.5,2.0]:
 p=d.add_paragraph(('Line spacing '+str(line)+' sample text. ')*20);p.paragraph_format.line_spacing=line;p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(18)
save(d,'spacing.docx')
d=fresh()
for name in ['List Bullet','List Bullet 2','List Bullet 3','List Number','List Number 2','List Number 3']:
 for i in range(2):d.add_paragraph(name+' item '+str(i+1),style=name)
save(d,'lists.docx')
d=fresh();t=d.add_table(rows=4,cols=3);t.style='Table Grid'
for r,row in enumerate(t.rows):
 for c,cell in enumerate(row.cells):cell.text=f'R{r+1}C{c+1}'
t.cell(0,0).merge(t.cell(0,1));t.cell(1,2).merge(t.cell(2,2));save(d,'table.docx')
d=fresh();p=d.add_paragraph('Mixed runs: ')
for label,attr in [('bold','bold'),('italic','italic'),('underline','underline'),('strike','strike')]:r=p.add_run(label+' ');setattr(r,attr,True)
r=p.add_run('red 18pt');r.font.color.rgb=RGBColor(200,30,40);r.font.size=Pt(18);save(d,'runs.docx')
d=fresh();s=d.sections[0];s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(3);s.right_margin=Cm(1.5);d.add_paragraph('Portrait A4 margins top2 bottom2 left3 right1.5 cm');d.add_page_break();d.add_paragraph('Explicit page break')
s=d.add_section(WD_SECTION_START.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width=Cm(29.7);s.page_height=Cm(21);s.left_margin=Cm(1.5);d.add_paragraph('Landscape A4 section');s.header.paragraphs[0].text='Synthetic header';s.footer.paragraphs[0].text='Synthetic footer';save(d,'sections.docx')
d=fresh()
for i in range(301):
 d.add_paragraph('Synthetic long document page '+str(i+1))
 if i<300:d.add_page_break()
save(d,'long.docx')
def pdf(name,count,mixed=False,encrypt=None,image=None):
 c=Canvas(str(out/name),pagesize=A4,encrypt=encrypt);c.setTitle('Synthetic viewer fixture');c.setAuthor('FreeOffice tester fixture generator')
 for i in range(count):
  size=landscape(A4) if mixed and i%2 else A4;c.setPageSize(size)
  if image:c.drawImage(str(image),30,30,width=size[0]-60,height=size[1]-60)
  else:
   c.drawString(40,size[1]-50,f'Page {i+1} / {count}: searchneedle repeated searchable text')
   c.drawString(40,size[1]-80,'The original file must remain read-only.')
  c.showPage()
 c.save()
pdf('text20.pdf',20);pdf('mixed-sizes.pdf',8,True);pdf('long301.pdf',301);pdf('password.pdf',2,encrypt=StandardEncryption('viewer-fixture',ownerPassword='synthetic-owner'))
im=Image.new('RGB',(600,800),'white');dr=ImageDraw.Draw(im);dr.text((40,50),'SCAN IMAGE ONLY - NO PDF TEXT LAYER',fill='black');im.save(out/'scan-image.png');pdf('scan-only.pdf',3,image=out/'scan-image.png')
inputs={'basic.csv':b'a,b,c\r\n00123,=SUM(A1),text\r\nlast,,tail','basic.tsv':b'a\tb\tc\n00123\t=literal\ttext','ragged.csv':b'a,b,c\n1,2\nx,y,z,tail','sep-lf.csv':b'sep=;\na;b\n1;2','sep-crlf.csv':b'sep=;\r\na;b\r\n1;2','nul.csv':b'a,b\n1,\x00','quote-error.csv':b'a,b\n1,"unterminated','utf16le.csv':b'\xff\xfe'+'a,b\r\n1,2'.encode('utf-16le'),'multiline.csv':b'a,b\n"line1\r\nline2","a""b"','empty.csv':b''}
for n,b in inputs.items():(out/n).write_bytes(b)
(out/'sample.txt').write_text('Synthetic text searchneedle\nSecond line\n'+'A long line. '*40,encoding='utf-8')
manifest={p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(out.iterdir()) if p.is_file()}
(out/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
(out/'README.md').write_text('''# Синтетические файлы для DroidDeck Tester

Файлы сгенерированы, но не открывались приложением и не испытаны. Это аналоги, а не оригинальные T12/T18/T21 или документы владельца.

DOCX: заголовки 1–3, отступы, четыре выравнивания, интервалы 1/1.5/2, стили списков 3 уровней, таблица 4 строки × 3 столбца с horizontal/vertical merge, смешанные runs, разделы A4/landscape с полями, 301 страница. Проверять against LibreOffice/reference и фиксировать реальные ограничения.

PDF: 20 текстовых страниц (searchneedle), scan-only без текстового слоя, mixed sizes, 301 страница, пароль viewer-fixture (публичная синтетическая фикстура, не секрет пользователя).

CSV/TSV: basic точные строки и ведущие нули, ragged длины 3/2/4, sep= LF/CRLF, отказ на NUL при сохранном UTF-16LE контроле, координаты ошибочной кавычки, многострочное поле, пустой ввод. Проверять значения и длины каждой сохранённой записи, а не только скриншоты.

manifest.json содержит входные размеры и SHA-256. Ни одному случаю не присвоен PASS.
''',encoding='utf-8')
print('Generated',len(manifest),'synthetic files; no tests executed:',out)
